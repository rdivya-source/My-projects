"""Core engine for the AI Resume Builder: parsing, matching, ATS audit, rewriting."""
import io
import re
from collections import Counter

# ---------------------------------------------------------------- skills
SKILLS = {
    "python": ["python"], "java": ["java"], "javascript": ["javascript", "js"],
    "typescript": ["typescript"], "c++": ["c++"], "sql": ["sql", "mysql", "postgresql", "postgres"],
    "html": ["html"], "css": ["css"], "react": ["react", "reactjs", "react.js"],
    "node.js": ["node.js", "nodejs", "node"], "django": ["django"], "flask": ["flask"],
    "streamlit": ["streamlit"], "fastapi": ["fastapi"], "rest apis": ["rest api", "rest apis", "restful"],
    "git": ["git", "github", "gitlab"], "docker": ["docker"], "kubernetes": ["kubernetes", "k8s"],
    "aws": ["aws", "amazon web services"], "azure": ["azure"], "linux": ["linux"],
    "machine learning": ["machine learning", "ml"], "deep learning": ["deep learning"],
    "nlp": ["nlp", "natural language processing"], "llm": ["llm", "llms", "large language model"],
    "pytorch": ["pytorch"], "tensorflow": ["tensorflow"], "scikit-learn": ["scikit-learn", "sklearn"],
    "pandas": ["pandas"], "numpy": ["numpy"], "data analysis": ["data analysis", "data analytics"],
    "data visualization": ["data visualization", "tableau", "power bi", "matplotlib"],
    "excel": ["excel"], "statistics": ["statistics", "statistical"],
    "agile": ["agile", "scrum"], "unit testing": ["unit testing", "unit tests", "pytest", "testing"],
    "ci/cd": ["ci/cd", "jenkins", "github actions"], "mongodb": ["mongodb"],
    "communication": ["communication"], "teamwork": ["teamwork", "collaboration", "team player"],
    "problem solving": ["problem solving", "problem-solving"], "leadership": ["leadership", "led a team"],
}

STOPWORDS = set("""a an the and or of to in for on with as at by from is are was were be been this that these those
we you your our their will can should must have has had it its into over per using use used etc looking seeking
candidate role job work working ability strong good excellent knowledge experience years year skills skill required
preferred plus team teams responsibilities requirements qualifications about who what""".split())

WEAK_TO_STRONG = {
    r"\bresponsible for\b": "Led", r"\bworked on\b": "Developed",
    r"\bin charge of\b": "Managed", r"\bduties included\b": "Delivered", r"\btried to\b": "Worked to",
    r"\bwas involved in\b": "Contributed to", r"\bmade\b": "Built",
}
ACTION_VERBS = set("""developed built designed led managed created implemented optimized improved delivered launched
analyzed automated engineered deployed reduced increased streamlined architected trained mentored coordinated
integrated tested achieved""".split())

DEGREES = [("phd", 4), ("doctorate", 4), ("master", 3), ("m.tech", 3), ("mba", 3), ("m.sc", 3), ("msc", 3),
           ("bachelor", 2), ("b.tech", 2), ("b.e", 2), ("b.sc", 2), ("bsc", 2), ("b.com", 2), ("diploma", 1)]


# ---------------------------------------------------------------- file reading
def extract_text(uploaded):
    """uploaded: object with .name and .read() (Streamlit UploadedFile) or a path string."""
    if isinstance(uploaded, str):
        name, data = uploaded, open(uploaded, "rb").read()
    else:
        name, data = uploaded.name, uploaded.read()
    name = name.lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        return "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
    if name.endswith(".docx"):
        import docx
        return "\n".join(p.text for p in docx.Document(io.BytesIO(data)).paragraphs)
    return data.decode("utf-8", errors="ignore")


# ---------------------------------------------------------------- parsing
def find_skills(text):
    t = " " + text.lower() + " "
    found = set()
    for canon, variants in SKILLS.items():
        for v in variants:
            if re.search(r"(?<![a-z0-9+#])" + re.escape(v) + r"(?![a-z0-9+#])", t):
                found.add(canon)
                break
    return found


def years_of_experience(text):
    m = re.findall(r"(\d+)\+?\s*(?:years|yrs)", text.lower())
    if m:
        return max(int(x) for x in m)
    spans = re.findall(r"((?:19|20)\d{2})\s*[-\u2013to]+\s*((?:19|20)\d{2}|present|current)", text.lower())
    total = 0
    for a, b in spans:
        b = 2026 if b in ("present", "current") else int(b)
        total += max(0, b - int(a))
    return total


def degree_level(text):
    t = text.lower()
    return max([lvl for k, lvl in DEGREES if k in t] or [0])


def parse_resume(text):
    return {
        "skills": find_skills(text),
        "years": years_of_experience(text),
        "degree": degree_level(text),
        "email": bool(re.search(r"[\w.+-]+@[\w-]+\.[\w.]+", text)),
        "phone": bool(re.search(r"(\+?\d[\d\s().-]{8,}\d)", text)),
        "words": len(text.split()),
    }


def parse_jd(text):
    """Split JD lines into required vs preferred skills."""
    req_lines, pref_lines, mode = [], [], "req"
    for line in text.splitlines():
        low = line.lower()
        if re.search(r"preferred|nice to have|good to have|bonus|plus\b", low):
            mode = "pref"
        elif re.search(r"required|requirements|must have|qualifications|responsibilities", low):
            mode = "req"
        (pref_lines if mode == "pref" else req_lines).append(line)
    required = find_skills("\n".join(req_lines))
    preferred = find_skills("\n".join(pref_lines)) - required
    words = [w for w in re.findall(r"[a-zA-Z][a-zA-Z+#.]{2,}", text.lower()) if w not in STOPWORDS]
    keywords = [w for w, _ in Counter(words).most_common(15)]
    return {"required": required, "preferred": preferred, "keywords": keywords,
            "years": years_of_experience(text), "degree": degree_level(text)}


# ---------------------------------------------------------------- matching (40/20/20/10/10)
def match_score(resume_text, jd_text):
    r, j = parse_resume(resume_text), parse_jd(jd_text)
    rt = resume_text.lower()

    def ratio(need, have):
        return 1.0 if not need else len(need & have) / len(need)

    req = 40 * ratio(j["required"], r["skills"])
    pref = 20 * ratio(j["preferred"], r["skills"])
    kw_hits = [k for k in j["keywords"] if k in rt]
    kw = 20 * (len(kw_hits) / len(j["keywords"]) if j["keywords"] else 1.0)
    exp = 10 * (1.0 if not j["years"] else min(1.0, r["years"] / j["years"]))
    edu = 10 * (1.0 if not j["degree"] else min(1.0, r["degree"] / j["degree"]))
    total = round(req + pref + kw + exp + edu)
    missing = sorted((j["required"] | j["preferred"]) - r["skills"])
    return {
        "total": total,
        "breakdown": [("Required Skills", round(req, 1), 40), ("Preferred Skills", round(pref, 1), 20),
                      ("Relevant Keywords", round(kw, 1), 20), ("Experience", round(exp, 1), 10),
                      ("Education", round(edu, 1), 10)],
        "matched": sorted((j["required"] | j["preferred"]) & r["skills"]),
        "missing": missing,
        "missing_required": sorted(j["required"] - r["skills"]),
        "keywords_missing": [k for k in j["keywords"] if k not in rt],
    }


# ---------------------------------------------------------------- ATS audit (15 checks)
def ats_audit(text, jd_text=""):
    t, low = text, text.lower()
    lines = [l.strip() for l in t.splitlines() if l.strip()]
    bullets = [l for l in lines if re.match(r"^[-\u2022*\u25aa\u25cf]", l)]
    body = [re.sub(r"^[-\u2022*\u25aa\u25cf]\s*", "", l) for l in bullets] or lines
    first_words = [b.split()[0].lower().strip(",.") for b in body if b.split()]
    verb_ratio = sum(w in ACTION_VERBS for w in first_words) / max(1, len(first_words))
    metrics = len(re.findall(r"\d+\s*%|\$\s*\d+|\b\d{2,}\b", t))
    words = len(t.split())
    kw = [k for k in parse_jd(jd_text)["keywords"]] if jd_text else []
    density = sum(k in low for k in kw) / len(kw) if kw else 1.0
    odd = len(re.findall(r"[^\x00-\x7F\u2022\u2013\u2014\u00b7\u25aa\u25cf]", t))
    checks = [
        ("Email address present", bool(re.search(r"[\w.+-]+@[\w-]+\.[\w.]+", t)), "Add a professional email at the top."),
        ("Phone number present", bool(re.search(r"\+?\d[\d\s().-]{8,}\d", t)), "Add a phone number."),
        ("LinkedIn / GitHub / portfolio link", bool(re.search(r"linkedin|github|portfolio|http", low)), "Add a LinkedIn or portfolio link."),
        ("Standard 'Experience' heading", bool(re.search(r"experience|employment|internship", low)), "Use a standard 'Experience' heading."),
        ("Standard 'Education' heading", "education" in low, "Add an 'Education' section."),
        ("Standard 'Skills' heading", "skills" in low, "Add a 'Skills' section."),
        ("Professional summary present", bool(re.search(r"summary|objective|profile", low)), "Add a 2-3 line professional summary."),
        ("Bullets start with action verbs", verb_ratio >= 0.4, "Start bullets with verbs like 'Built', 'Led', 'Optimized'."),
        ("Quantifiable achievements", metrics >= 3, "Add numbers (%, $, counts) to show impact."),
        ("Keyword coverage vs job description", density >= 0.5, "Work more job-description keywords into your resume."),
        ("Length is 250-900 words", 250 <= words <= 900, "Aim for one to two pages (250-900 words)."),
        ("Uses bullet points", len(bullets) >= 4, "Use bullet points for achievements."),
        ("No unusual symbols / graphics", odd < 5, "Remove icons/special characters that ATS parsers can't read."),
        ("No first-person pronouns", not re.search(r"\b(i|my|me)\b", low), "Drop 'I' / 'my' - resumes use implied first person."),
        ("Consistent date format", len({("mon" if re.search(r"[a-z]{3}", d) else "num") for d in re.findall(r"\b(?:[A-Za-z]{3,9}\s+)?\d{4}\b|\b\d{1,2}/\d{4}\b", t)}) <= 1,
         "Use one date format throughout."),
    ]
    return [{"check": c, "pass": bool(p), "fix": f} for c, p, f in checks]


# ---------------------------------------------------------------- rewriting
def demo_rewrite(resume_text, missing_skills):
    """Offline DemoModel: deterministic rule-based improvement."""
    out, changes = [], []
    for line in resume_text.splitlines():
        new = line
        for pat, rep in WEAK_TO_STRONG.items():
            if re.search(pat, new, flags=re.I):
                new = re.sub(pat, rep, new, count=1, flags=re.I)
                new = new[:1].upper() + new[1:] if new.strip() else new
        if new != line:
            changes.append((line.strip(), new.strip()))
        if re.match(r"^\s*[-\u2022*]", new) and not re.search(r"\d", new):
            new += "  [add a metric, e.g. 'improved X by 25%']"
        out.append(new)
    text = "\n".join(out)
    if missing_skills:
        text += "\n\nSUGGESTED SKILLS TO ADD (only if you genuinely have them): " + ", ".join(missing_skills)
    return text, changes


def llm_rewrite(resume_text, jd_text, api_key, model="gpt-4o-mini", few_shot=True, evidence_text=""):
    """Optional LLM mode (OpenAI). Falls back to demo_rewrite on any error."""
    try:
        from openai import OpenAI
        shots = ("Example: 'Worked on website' -> 'Developed a responsive website used by 5,000+ users, "
                 "cutting load time by 30%.'\n") if few_shot else ""
        prompt = (f"Rewrite this resume to better match the job description. Keep all facts truthful; "
                  f"do NOT invent experience. Use strong action verbs and ATS-friendly plain formatting. "
                  f"Only claim skills that have supporting evidence below.\n{shots}\n"
                  f"RETRIEVED EVIDENCE (from the resume):\n{evidence_text}\n\n"
                  f"JOB DESCRIPTION:\n{jd_text}\n\nRESUME:\n{resume_text}")
        r = OpenAI(api_key=api_key).chat.completions.create(
            model=model, messages=[{"role": "user", "content": prompt}], temperature=0.3)
        return r.choices[0].message.content, []
    except Exception as e:  # graceful fallback, as described in the report
        text, ch = demo_rewrite(resume_text, [])
        return f"[LLM unavailable ({type(e).__name__}) - fell back to offline DemoModel]\n\n{text}", ch


def to_docx_bytes(text):
    import docx
    d = docx.Document()
    for line in text.splitlines():
        d.add_paragraph(line)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
